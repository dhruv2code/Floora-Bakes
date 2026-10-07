from datetime import date, time
from io import BytesIO

import pandas as pd
from fastapi import status
from openpyxl import load_workbook

from app.models.sales import Sales, UploadBatch


def test_report_export_filters_sales(api_client, database):
    batch = UploadBatch(
        filename="sales.xlsx",
        total_records=2,
        valid_records=2,
        invalid_records=0,
        duplicate_records=0,
        mapping="{}",
    )
    database.add(batch)
    database.flush()
    database.add_all([
        Sales(
            sale_id="S1",
            item_name="Cake",
            sales_date=date(2026, 4, 1),
            day="Wednesday",
            price=120,
            sell_time=time(17, 30),
            upload_batch_id=batch.id,
        ),
        Sales(
            sale_id="S2",
            item_name="Cookie",
            sales_date=date(2026, 4, 2),
            day="Thursday",
            price=20,
            sell_time=time(10, 0),
            upload_batch_id=batch.id,
        ),
    ])
    database.commit()

    response = api_client.get(
        "/api/reports/export",
        params={"format": "csv", "start_date": "2026-04-02", "end_date": "2026-04-02"},
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.headers["content-type"].startswith("text/csv")
    assert "S2,Cookie,2026-04-02,Thursday,20.0,10:00" in response.text
    assert "S1" not in response.text


def test_upload_and_delete_batch(api_client, database):
    frame = pd.DataFrame([
        {
            "Sale ID": "S1",
            "Item Name": "Cake",
            "Sales Date": "2026-04-01",
            "Day": "Wednesday",
            "Price": "120",
            "Sell Time": "17:30",
        }
    ])
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
    buffer.seek(0)

    response = api_client.post(
        "/api/upload",
        files=[
            (
                "file",
                ("sales.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ),
            (
                "mapping",
                (
                    None,
                    '{"sale_id":"Sale ID","item_name":"Item Name","sales_date":"Sales Date","day":"Day","price":"Price","sell_time":"Sell Time"}',
                ),
            ),
        ],
    )

    assert response.status_code == status.HTTP_201_CREATED, response.text
    import_result = response.json()
    assert import_result["imported_records"] == 1

    history = api_client.get("/api/upload/history")
    assert history.status_code == status.HTTP_200_OK
    assert history.json()[0]["id"] == import_result["batch_id"]

    delete_response = api_client.delete(f"/api/upload/batches/{import_result['batch_id']}")
    assert delete_response.status_code == status.HTTP_204_NO_CONTENT
    assert database.query(Sales).count() == 0
    assert database.query(UploadBatch).count() == 0


def test_xlsx_report_contains_summary_and_sales_sheets(api_client, database):
    batch = UploadBatch(
        filename="sales.xlsx",
        total_records=1,
        valid_records=1,
        invalid_records=0,
        duplicate_records=0,
        mapping="{}",
    )
    database.add(batch)
    database.flush()
    database.add(Sales(
        sale_id="S1",
        item_name="Cake",
        sales_date=date(2026, 4, 1),
        day="Wednesday",
        price=120,
        sell_time=time(17, 30),
        upload_batch_id=batch.id,
    ))
    database.commit()

    response = api_client.get("/api/reports/export", params={"format": "xlsx"})

    assert response.status_code == status.HTTP_200_OK
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Sales", "Top Products", "Daily Analysis", "Time Analysis"]
